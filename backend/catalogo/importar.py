"""Importación y exportación de productos en Excel (.xlsx) o CSV."""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Categoria, Producto
from .services import mover_stock, siguiente_codigo_interno

COLUMNAS = ['codigo_barras', 'nombre', 'categoria', 'precio_costo', 'precio_venta',
            'unidad_medida', 'stock', 'stock_minimo']
ALIAS = {
    'codigo': 'codigo_barras', 'código': 'codigo_barras', 'codigo de barras': 'codigo_barras',
    'código de barras': 'codigo_barras', 'producto': 'nombre', 'descripcion': 'nombre',
    'costo': 'precio_costo', 'precio costo': 'precio_costo', 'precio': 'precio_venta',
    'precio venta': 'precio_venta', 'unidad': 'unidad_medida', 'cantidad': 'stock',
    'existencia': 'stock', 'existencias': 'stock', 'minimo': 'stock_minimo', 'mínimo': 'stock_minimo',
    'stock minimo': 'stock_minimo', 'stock mínimo': 'stock_minimo',
}
UNIDADES = {c[0] for c in Producto.UNIDAD_CHOICES}


def _num(v, defecto='0'):
    if v is None or str(v).strip() == '':
        return Decimal(defecto)
    s = str(v).strip().replace('$', '').replace(' ', '')
    # 1.500,50 -> 1500.50 ; 1,500.50 -> 1500.50 ; 1500 -> 1500
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.') if len(s.split(',')[-1]) != 3 else s.replace(',', '')
    elif s.count('.') > 1 or (s.count('.') == 1 and len(s.split('.')[-1]) == 3 and s.split('.')[0] not in ('0', '-0')):
        s = s.replace('.', '')  # 1.200 -> 1200 (separador de miles colombiano)
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f'"{v}" no es un número válido')


def _filas(archivo):
    nombre = (archivo.name or '').lower()
    if nombre.endswith('.csv'):
        texto = archivo.read().decode('utf-8-sig', errors='replace')
        dialect = csv.Sniffer().sniff(texto[:2000], delimiters=',;\t')
        filas = list(csv.reader(io.StringIO(texto), dialect))
    else:
        wb = load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    if not filas:
        return [], []
    enc = [ALIAS.get(str(c or '').strip().lower(), str(c or '').strip().lower()) for c in filas[0]]
    return enc, filas[1:]


@transaction.atomic
def importar(archivo, usuario, actualizar_existentes=True):
    enc, filas = _filas(archivo)
    if 'nombre' not in enc or 'precio_venta' not in enc:
        raise ValueError('El archivo debe tener al menos las columnas "nombre" y "precio_venta".')
    idx = {c: enc.index(c) for c in COLUMNAS if c in enc}
    creados = actualizados = 0
    errores = []
    cats = {c.nombre.lower(): c for c in Categoria.objects.all()}

    for n, fila in enumerate(filas, start=2):
        def val(col):
            i = idx.get(col)
            return fila[i] if i is not None and i < len(fila) else None

        nombre = str(val('nombre') or '').strip()
        if not nombre:
            continue
        try:
            with transaction.atomic():
                cat_nombre = str(val('categoria') or 'General').strip() or 'General'
                cat = cats.get(cat_nombre.lower())
                if not cat:
                    cat = Categoria.objects.create(nombre=cat_nombre[:80])
                    cats[cat.nombre.lower()] = cat
                codigo = str(val('codigo_barras') or '').strip()
                if codigo.endswith('.0'):
                    codigo = codigo[:-2]
                unidad = str(val('unidad_medida') or 'UNIDAD').strip().upper()
                unidad = unidad if unidad in UNIDADES else 'UNIDAD'
                datos = {
                    'nombre': nombre[:200], 'categoria': cat,
                    'precio_costo': _num(val('precio_costo')),
                    'precio_venta': _num(val('precio_venta')),
                    'unidad_medida': unidad,
                    'stock_minimo': _num(val('stock_minimo')),
                }
                existente = Producto.objects.filter(codigo_barras=codigo).first() if codigo else \
                    Producto.objects.filter(nombre__iexact=nombre).first()
                stock = val('stock')
                if existente:
                    if not actualizar_existentes:
                        continue
                    for k, v in datos.items():
                        setattr(existente, k, v)
                    existente.activo = True
                    existente.save()
                    if stock not in (None, ''):
                        mover_stock(existente, 'CONTEO', _num(stock), usuario, motivo='Importación de productos',
                                    referencia_tipo='IMPORTACION')
                    actualizados += 1
                else:
                    p = Producto.objects.create(codigo_barras=codigo or siguiente_codigo_interno(), **datos)
                    if stock not in (None, '') and _num(stock) > 0:
                        mover_stock(p, 'INICIAL', _num(stock), usuario, motivo='Importación de productos',
                                    referencia_tipo='IMPORTACION')
                    creados += 1
        except Exception as e:
            errores.append(f'Fila {n} ({nombre}): {e}')
    return {'creados': creados, 'actualizados': actualizados, 'errores': errores[:50],
            'total_errores': len(errores)}


def _libro(titulo):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1570EF')
    for col, ancho in zip('ABCDEFGH', (18, 40, 18, 14, 14, 14, 10, 12)):
        ws.column_dimensions[col].width = ancho
    return wb, ws


def plantilla() -> bytes:
    wb, ws = _libro('Productos')
    ws.append(['7702004003508', 'Arroz Diana 500 g', 'Granos', 1900, 2500, 'UNIDAD', 24, 6])
    ws.append(['', 'Tomate chonto', 'Fruver', 0, 0, 'KG', '', ''])
    ws.append(['7702129000017', 'Gaseosa 400 ml', 'Bebidas', 1600, 2200, 'BOTELLA', 48, 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar(qs) -> bytes:
    wb, ws = _libro('Productos')
    for p in qs.select_related('categoria'):
        ws.append([p.codigo_barras or '', p.nombre, p.categoria.nombre, float(p.precio_costo),
                   float(p.precio_venta), p.unidad_medida, float(p.stock), float(p.stock_minimo)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
