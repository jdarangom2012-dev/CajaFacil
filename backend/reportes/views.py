import io
from datetime import timedelta
from decimal import Decimal

from django.conf import settings
from django.db.models import Count, DecimalField, F, Q, Sum
from django.db.models.functions import ExtractHour, TruncDate
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from catalogo.models import Producto
from core import licencia, sync
from core.models import Configuracion
from core.permisos import GESTION, Roles
from core.utils import a_float, hoy, parse_fecha, rango_fechas
from creditos.models import AbonoCredito, Credito
from ventas.models import Venta, VentaDetalle, VentaPago

DEC = DecimalField(max_digits=14, decimal_places=2)


def _periodo(request, dias_defecto=0):
    h = hoy()
    desde = parse_fecha(request.query_params.get('fecha_desde'), h - timedelta(days=dias_defecto))
    hasta = parse_fecha(request.query_params.get('fecha_hasta'), h)
    if desde > hasta:
        desde, hasta = hasta, desde
    return desde, hasta


def _ventas(desde, hasta):
    ini, fin = rango_fechas(desde, hasta)
    return Venta.objects.filter(creado__gte=ini, creado__lt=fin, estado='COMPLETADA')


def _totales(qs):
    agg = qs.aggregate(total=Sum('total'), costo=Sum('costo_total'), n=Count('id'), desc=Sum('descuento'))
    total = a_float(agg['total'])
    n = agg['n'] or 0
    return {
        'num_ventas': n,
        'total_ventas': total,
        'utilidad': total - a_float(agg['costo']),
        'descuentos': a_float(agg['desc']),
        'ticket_promedio': round(total / n, 0) if n else 0,
    }


def _por_forma_pago(qs):
    return [{'nombre': r['forma_pago__nombre'], 'tipo': r['forma_pago__tipo'], 'color': r['forma_pago__color'],
             'total': a_float(r['t']), 'cantidad': r['n']}
            for r in VentaPago.objects.filter(venta__in=qs)
            .values('forma_pago__nombre', 'forma_pago__tipo', 'forma_pago__color')
            .annotate(t=Sum('monto'), n=Count('id')).order_by('-t')]


def _top_productos(qs, limite=10):
    return [{'producto_id': r['producto'], 'nombre': r['nombre_producto'], 'cantidad': a_float(r['cant']),
             'total': a_float(r['t']), 'utilidad': a_float(r['t']) - a_float(r['costo'])}
            for r in VentaDetalle.objects.filter(venta__in=qs)
            .values('producto', 'nombre_producto')
            .annotate(cant=Sum('cantidad'), t=Sum('subtotal'), costo=Sum(F('costo_unit') * F('cantidad'), output_field=DEC))
            .order_by('-t')[:limite]]


def _por_categoria(qs):
    filas = list(VentaDetalle.objects.filter(venta__in=qs)
                 .values('producto__categoria__nombre', 'producto__categoria__color')
                 .annotate(cant=Sum('cantidad'), t=Sum('subtotal'),
                           costo=Sum(F('costo_unit') * F('cantidad'), output_field=DEC))
                 .order_by('-t'))
    total = sum(a_float(r['t']) for r in filas) or 1
    return [{'categoria': r['producto__categoria__nombre'], 'color': r['producto__categoria__color'],
             'cantidad': a_float(r['cant']), 'total': a_float(r['t']),
             'utilidad': a_float(r['t']) - a_float(r['costo']),
             'porcentaje': round(a_float(r['t']) / total * 100, 1)} for r in filas]


def _por_dia(qs):
    return [{'fecha': r['d'], 'total': a_float(r['t']), 'cantidad': r['n'],
             'utilidad': a_float(r['t']) - a_float(r['c'])}
            for r in qs.annotate(d=TruncDate('creado')).values('d')
            .annotate(t=Sum('total'), n=Count('id'), c=Sum('costo_total')).order_by('d')]


def _por_cajero(qs):
    return [{'usuario': r['usuario__nombre_completo'] or r['usuario__username'], 'total': a_float(r['t']),
             'cantidad': r['n']}
            for r in qs.values('usuario__nombre_completo', 'usuario__username')
            .annotate(t=Sum('total'), n=Count('id')).order_by('-t')]


class DashboardView(APIView):
    def get(self, request):
        h = hoy()
        qs_hoy = _ventas(h, h)
        qs_ayer = _ventas(h - timedelta(days=1), h - timedelta(days=1))
        semana = _ventas(h - timedelta(days=6), h)
        ini, fin = rango_fechas(h, h)
        por_hora = {r['h']: a_float(r['t']) for r in qs_hoy.annotate(h=ExtractHour('creado'))
                    .values('h').annotate(t=Sum('total'))}
        bajo = Producto.objects.filter(activo=True, categoria__controla_stock=True, stock__lte=F('stock_minimo'))
        pend = Credito.objects.filter(estado='PENDIENTE')
        datos = {
            'fecha': h,
            'hoy': _totales(qs_hoy),
            'ayer': _totales(qs_ayer),
            'anuladas_hoy': Venta.objects.filter(creado__gte=ini, creado__lt=fin, estado='ANULADA').count(),
            'por_forma_pago': _por_forma_pago(qs_hoy),
            'por_hora': [{'hora': hh, 'total': por_hora.get(hh, 0)} for hh in range(6, 23)],
            'top_productos': _top_productos(qs_hoy, 5),
            'semana': _por_dia(semana),
            'ultimas_ventas': [
                {'id': v.id, 'numero': v.numero, 'creado': v.creado, 'total': v.total,
                 'usuario': v.usuario.nombre, 'cliente': v.cliente.nombre if v.cliente else None}
                for v in qs_hoy.select_related('usuario', 'cliente').order_by('-creado')[:8]],
            'stock_bajo': {
                'total': bajo.count(),
                'productos': list(bajo.order_by('stock').values('id', 'nombre', 'stock', 'stock_minimo',
                                                               'unidad_medida')[:8]),
            },
            'fiado': {
                'cartera': pend.aggregate(s=Sum(F('total') - F('abonado'), output_field=DEC))['s'] or 0,
                'vencidos': pend.filter(vence__lt=h).count(),
                'abonos_hoy': AbonoCredito.objects.filter(creado__gte=ini, creado__lt=fin).aggregate(
                    s=Sum('monto'))['s'] or 0,
            },
        }
        if not Roles(*GESTION)().has_permission(request, self):
            # El cajero no ve utilidades
            for k in ('hoy', 'ayer'):
                datos[k].pop('utilidad', None)
            for d in datos['semana']:
                d.pop('utilidad', None)
        return Response(datos)


class ResumenPeriodoView(APIView):
    permission_classes = [IsAuthenticated, Roles(*GESTION)]

    def get(self, request):
        desde, hasta = _periodo(request, 6)
        qs = _ventas(desde, hasta)
        ini, fin = rango_fechas(desde, hasta)
        anuladas = Venta.objects.filter(creado__gte=ini, creado__lt=fin, estado='ANULADA')
        return Response({
            'desde': desde, 'hasta': hasta,
            'totales': _totales(qs),
            'anuladas': {'cantidad': anuladas.count(), 'total': a_float(anuladas.aggregate(t=Sum('total'))['t'])},
            'por_dia': _por_dia(qs),
            'por_forma_pago': _por_forma_pago(qs),
            'por_categoria': _por_categoria(qs),
            'top_productos': _top_productos(qs, 15),
            'por_cajero': _por_cajero(qs),
            'fiado': {
                'otorgado': a_float(Credito.objects.filter(creado__gte=ini, creado__lt=fin)
                                    .exclude(estado='ANULADO').aggregate(t=Sum('total'))['t']),
                'abonado': a_float(AbonoCredito.objects.filter(creado__gte=ini, creado__lt=fin)
                                   .aggregate(t=Sum('monto'))['t']),
            },
        })


class VentasProductoView(APIView):
    permission_classes = [IsAuthenticated, Roles(*GESTION)]

    def get(self, request):
        desde, hasta = _periodo(request, 29)
        qs = _ventas(desde, hasta)
        filas = (VentaDetalle.objects.filter(venta__in=qs)
                 .values('producto', 'nombre_producto', 'producto__categoria__nombre')
                 .annotate(cant=Sum('cantidad'), t=Sum('subtotal'), n=Count('venta', distinct=True),
                           costo=Sum(F('costo_unit') * F('cantidad'), output_field=DEC))
                 .order_by('-t'))
        q = request.query_params.get('q')
        if q:
            filas = filas.filter(nombre_producto__icontains=q)
        if request.query_params.get('categoria'):
            filas = filas.filter(producto__categoria_id=request.query_params['categoria'])
        data = [{'producto_id': r['producto'], 'nombre': r['nombre_producto'],
                 'categoria': r['producto__categoria__nombre'], 'cantidad': a_float(r['cant']),
                 'total': a_float(r['t']), 'ventas': r['n'], 'utilidad': a_float(r['t']) - a_float(r['costo'])}
                for r in filas[:1000]]
        return Response({'desde': desde, 'hasta': hasta, 'productos': data,
                         'total': sum(x['total'] for x in data)})


class PedidoSugeridoView(APIView):
    """
    Predicción liviana de agotamiento (reemplaza a Prophet, que es muy pesado para un instalador):
    promedio de venta diaria de los últimos N días → días que alcanza el stock y cuánto pedir
    para cubrir X días.
    """
    permission_classes = [IsAuthenticated, Roles(*GESTION, 'ALMACENISTA')]

    def get(self, request):
        dias_historia = int(request.query_params.get('dias_historia', 28))
        dias_cobertura = int(request.query_params.get('dias_cobertura', 7))
        h = hoy()
        qs = _ventas(h - timedelta(days=dias_historia - 1), h)
        vendidos = {r['producto']: a_float(r['c']) for r in
                    VentaDetalle.objects.filter(venta__in=qs).values('producto').annotate(c=Sum('cantidad'))}
        resultado = []
        for p in Producto.objects.filter(activo=True, categoria__controla_stock=True).select_related('categoria'):
            prom = vendidos.get(p.id, 0) / dias_historia
            stock = float(p.stock)
            if prom <= 0 and stock > float(p.stock_minimo):
                continue
            dias_restantes = round(stock / prom, 1) if prom > 0 else None
            objetivo = max(prom * dias_cobertura, float(p.stock_minimo))
            sugerido = max(0, objetivo - stock)
            if p.unidad_medida in ('UNIDAD', 'PAQUETE', 'CAJA', 'BOTELLA', 'BOLSA'):
                sugerido = float(int(sugerido + 0.999))
            if sugerido <= 0 and (dias_restantes is None or dias_restantes > dias_cobertura):
                continue
            resultado.append({
                'producto_id': p.id, 'nombre': p.nombre, 'categoria': p.categoria.nombre,
                'unidad': p.unidad_medida, 'stock': stock, 'stock_minimo': float(p.stock_minimo),
                'venta_diaria': round(prom, 2), 'dias_restantes': dias_restantes,
                'sugerido': round(sugerido, 2), 'costo_estimado': round(sugerido * float(p.precio_costo)),
                'urgencia': 'CRITICO' if (dias_restantes is not None and dias_restantes <= 2) or stock <= 0
                else 'PRONTO' if dias_restantes is not None and dias_restantes <= dias_cobertura else 'NORMAL',
            })
        orden = {'CRITICO': 0, 'PRONTO': 1, 'NORMAL': 2}
        resultado.sort(key=lambda x: (orden[x['urgencia']], x['dias_restantes'] if x['dias_restantes'] is not None else 9999))
        return Response({'dias_historia': dias_historia, 'dias_cobertura': dias_cobertura,
                         'total_costo': sum(x['costo_estimado'] for x in resultado), 'productos': resultado})


class AsistenteView(APIView):
    """Asistente IA (plan nube): la pregunta se resuelve en la nube de NautilusTech con datos resumidos."""
    permission_classes = [IsAuthenticated, Roles(*GESTION)]

    def post(self, request):
        import requests
        est = licencia.estado_actual()
        if not est.nube_activa:
            return Response({'error': 'El asistente con IA está incluido en el plan Nube. '
                                      'Actívalo para preguntarle a tu tienda.'}, status=402)
        pregunta = (request.data.get('mensaje') or '').strip()[:500]
        if not pregunta:
            return Response({'error': 'Escribe tu pregunta.'}, status=400)
        h = hoy()
        contexto = {
            'negocio': Configuracion.obtener().nombre_negocio,
            'hoy': _totales(_ventas(h, h)),
            'ultimos_30_dias': _totales(_ventas(h - timedelta(days=29), h)),
            'top_productos_30d': _top_productos(_ventas(h - timedelta(days=29), h), 15),
            'por_categoria_30d': _por_categoria(_ventas(h - timedelta(days=29), h)),
        }
        try:
            r = requests.post(f'{settings.NUBE_URL}/api/ia/preguntar',
                              json={'pregunta': pregunta, 'contexto': contexto, 'licencia_id': est.licencia_id,
                                    'equipo_id': est.equipo}, headers=sync._cabeceras(est), timeout=40)
            if r.status_code in (402, 503):
                return Response({'error': sync._mensaje_error(r)}, status=r.status_code)
            r.raise_for_status()
            return Response({'respuesta': r.json().get('respuesta', '')})
        except Exception:
            return Response({'error': 'No hay conexión con el asistente. Revisa el internet e intenta de nuevo.'},
                            status=503)


def _xlsx(titulo, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:30]
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1570EF')
    for f in filas:
        ws.append(f)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(45, max(len(str(c.value or '')) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ExportarVentasView(APIView):
    permission_classes = [IsAuthenticated, Roles(*GESTION)]

    def get(self, request):
        desde, hasta = _periodo(request, 0)
        ini, fin = rango_fechas(desde, hasta)
        tipo = request.query_params.get('tipo', 'ventas')
        if tipo == 'detalle':
            qs = (VentaDetalle.objects.filter(venta__creado__gte=ini, venta__creado__lt=fin)
                  .select_related('venta', 'producto__categoria').order_by('venta__numero'))
            filas = [[d.venta.numero, timezone.localtime(d.venta.creado).strftime('%Y-%m-%d %H:%M'),
                      d.venta.estado, d.nombre_producto, d.producto.categoria.nombre, float(d.cantidad),
                      float(d.precio_unit), float(d.descuento), float(d.subtotal),
                      float(d.subtotal - d.costo_unit * d.cantidad)] for d in qs]
            contenido = _xlsx('Detalle', ['Venta', 'Fecha', 'Estado', 'Producto', 'Categoría', 'Cantidad',
                                          'Precio', 'Descuento', 'Subtotal', 'Utilidad'], filas)
        else:
            qs = (Venta.objects.filter(creado__gte=ini, creado__lt=fin)
                  .select_related('usuario', 'cliente').prefetch_related('pagos__forma_pago').order_by('numero'))
            filas = [[v.numero, timezone.localtime(v.creado).strftime('%Y-%m-%d %H:%M'), v.terminal,
                      v.usuario.nombre, v.cliente.nombre if v.cliente else '', v.estado, float(v.subtotal),
                      float(v.descuento), float(v.total), float(v.total - v.costo_total),
                      ', '.join(f'{p.forma_pago.nombre} {float(p.monto):,.0f}' for p in v.pagos.all()),
                      v.motivo_anulacion] for v in qs]
            contenido = _xlsx('Ventas', ['Número', 'Fecha', 'Caja', 'Cajero', 'Cliente', 'Estado', 'Subtotal',
                                         'Descuento', 'Total', 'Utilidad', 'Pagos', 'Motivo anulación'], filas)
        r = HttpResponse(contenido, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r['Content-Disposition'] = f'attachment; filename="{tipo}_{desde}_{hasta}.xlsx"'
        return r
